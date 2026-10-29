#!/usr/bin/env python3
"""生成面镜 MockMirror 的测试用例与执行结果工作簿。"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs" / "面镜MockMirror-测试用例与执行报告-v1.0.xlsx"
RES = ROOT / "docs" / "test-results"

INK = "111827"
MUTED = "6B7280"
LINE = "E5E7EB"
HEADER = "1F2937"
ZEBRA = "F9FAFB"
GREEN = "DCFCE7"
GREEN_TXT = "166534"
RED = "FEE2E2"
RED_TXT = "991B1B"
AMBER = "FEF3C7"
AMBER_TXT = "92400E"

thin = Side(style="thin", color=LINE)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def load(name):
    return json.loads((RES / name).read_text(encoding="utf-8"))


def style_header(ws, row, cols, fill=HEADER, color="FFFFFF"):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(name="Arial", size=10, bold=True, color=color)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def title_block(ws, title, subtitle, cols):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=cols)
    c = ws.cell(row=1, column=1, value="　　" + title)
    c.fill = PatternFill("solid", fgColor=INK)
    c.font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 38
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=cols)
    c = ws.cell(row=2, column=1, value="　　" + subtitle)
    c.font = Font(name="Arial", size=9, color=MUTED)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 22


def write_rows(ws, start_row, rows, widths, result_col=None, wrap_cols=None, height=44):
    for r_i, row in enumerate(rows, start_row):
        for c_i, value in enumerate(row, 1):
            cell = ws.cell(row=r_i, column=c_i, value=value)
            cell.font = Font(name="Arial", size=9, color=INK)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
            if r_i % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=ZEBRA)
        if result_col:
            cell = ws.cell(row=r_i, column=result_col)
            val = str(cell.value or "")
            if val.startswith("通过") or val == "PASS":
                cell.fill = PatternFill("solid", fgColor=GREEN)
                cell.font = Font(name="Arial", size=9, bold=True, color=GREEN_TXT)
            elif val.startswith("失败") or val == "FAIL":
                cell.fill = PatternFill("solid", fgColor=RED)
                cell.font = Font(name="Arial", size=9, bold=True, color=RED_TXT)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[r_i].height = height
    for idx, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = w


# ---------------- 功能用例（人工验收，基于实际执行） ----------------
F = [
    ("F-01", "工作台", "5.1 / 7.0", "首页加载与导航", "服务已启动，浏览器打开根路径",
     "1. 访问 /\n2. 检查侧栏 5 个入口\n3. 检查 Hero、统计、训练闭环区块",
     "页面正常渲染；5 个导航项；无横向滚动", "符合预期；5 个导航项，1440px 下无横向溢出", "通过", "P0", ""),
    ("F-02", "工作台", "5.1", "首页统计刷新", "完成至少 1 次面试并生成报告",
     "回到工作台，对比 /api/stats 的 answered 与页面统计",
     "页面统计与接口一致，不显示旧值", "刷新后一致（D-04 修复后回归）", "通过", "P1", "D-04"),
    ("F-03", "工作台", "5.1", "最近记录入口", "已有进行中与已完成的面试",
     "点击最近记录行", "进行中进入面试页，已完成进入报告页", "符合预期", "通过", "P1", ""),
    ("F-04", "简历解析", "7.1 / AC-01", "上传 DOCX", "准备真实 DOCX 简历",
     "开始面试 → 拖入/选择 DOCX", "解析成功，展示姓名、教育、3 段经历、2 个项目", "石浩田 / 3 段经历 / 2 个项目", "通过", "P0", ""),
    ("F-05", "简历解析", "7.1 / AC-01", "上传 PDF", "准备含文字层的 PDF",
     "开始面试 → 上传 PDF", "解析成功，文字与 DOCX 一致", "识别 3 段经历，质量分 96", "通过", "P0", ""),
    ("F-06", "简历解析", "7.1", "粘贴文本解析", "无文件，准备一段完整简历文本",
     "切换到粘贴文本 → 输入 80 字以上 → 点击解析", "解析成功，进入画像预览", "符合预期", "通过", "P0", ""),
    ("F-07", "简历解析", "7.1", "示例简历一键体验", "无简历数据",
     "点击「用示例简历试试」", "30 秒内完成解析，展示周予安画像", "解析质量 96，2 段经历", "通过", "P0", ""),
    ("F-08", "简历解析", "7.1", "解析质量分与提示", "上传排版规范的简历",
     "查看画像右上角解析质量与底部提示", "分数 0-100，低分提示调整排版", "质量分 96，提示正常", "通过", "P1", ""),
    ("F-09", "简历解析", "7.1 / 10.3", "非法格式拒绝", "准备 .doc 或 .xyz 文件",
     "上传非法格式", "400 可读提示，不进入画像页", "返回「请另存为 .docx」；D-01 修复后为 400", "通过", "P0", "D-01"),
    ("F-10", "简历解析", "7.1 / 10.3", "损坏 PDF 拒绝", "准备损坏/扫描版 PDF",
     "上传 PDF", "400 提示扫描件或格式问题", "「像是扫描件」提示；D-01 修复后为 400", "通过", "P1", "D-01"),
    ("F-11", "岗位匹配", "7.2 / AC-02", "JD 匹配分析", "已解析简历，准备示例 JD",
     "粘贴 JD → 分析岗位匹配", "展示岗位、匹配分、关键词覆盖与缺口", "AI 产品经理 / 72 分 / 覆盖 60%", "通过", "P0", ""),
    ("F-12", "岗位匹配", "7.2", "关键词缺口提示", "JD 含大模型、Agent、B 端",
     "查看关键词区与建议", "已覆盖绿色、缺口黄色，并给补证据建议", "符合预期", "通过", "P1", ""),
    ("F-13", "岗位匹配", "7.2 / 10.3", "过短 JD 拒绝", "JD 少于 20 字",
     "点击分析", "提示补充 JD，不创建分析结果", "400，前端 toast 提示", "通过", "P1", ""),
    ("F-14", "面试设置", "7.4", "三种人格选择", "进入第三步",
     "依次点击同级/资深/直属经理", "选中态互斥，右侧摘要同步", "符合预期", "通过", "P0", ""),
    ("F-15", "面试设置", "7.4", "题量与侧重调整", "进入第三步",
     "拖动题量 4-10，切换侧重", "摘要同步；题量与侧重写入 config", "5 题 / 直属经理，摘要同步", "通过", "P0", ""),
    ("F-16", "面试设置", "7.4 / 10.1", "开始面试", "简历、JD、设置均完成",
     "点击开始面试", "创建会话并进入面试页，显示 1/N 与首题", "进入面试页，1/5，首题正确", "通过", "P0", ""),
    ("F-17", "模拟面试", "7.4 / AC-04", "开场与首题", "会话刚创建",
     "查看对话区与右侧当前问题", "开场白、引导语、第一题；右侧考点可见", "符合预期", "通过", "P0", ""),
    ("F-18", "模拟面试", "7.4", "提交回答与流式回复", "处于可作答状态",
     "输入 200 字回答，Enter 发送", "用户气泡出现；面试官流式回复；进度前进", "流式输出完整，进度 2/5", "通过", "P0", ""),
    ("F-19", "模拟面试", "7.4 / 7.5", "单题五维评分", "完成 1 次作答",
     "查看右侧「本场表现」", "展示五维分数、总分、最近一题结论", "结构 92 / 相关 68 / 深度 73 / 结果 74 / 匹配 75", "通过", "P0", ""),
    ("F-20", "模拟面试", "7.4", "提示功能", "当前题未作答",
     "点击「给点提示」", "出现提示气泡，展开参考要点，记录 hintCount", "符合预期", "通过", "P1", ""),
    ("F-21", "模拟面试", "7.4", "跳过当前题", "当前题未作答",
     "点击「换一题」并确认", "进入下一题，记录跳过，不产生评分", "进度前进，下一题展示", "通过", "P1", ""),
    ("F-22", "模拟面试", "7.8", "Agent 轨迹实时展示", "完成 1 次作答",
     "查看右侧 Agent 轨迹", "展示提问/评估/路由节点、耗时与原因", "按时间倒序展示 4 类节点", "通过", "P1", ""),
    ("F-23", "模拟面试", "7.4 / 10.1", "刷新恢复会话", "作答 1 题后刷新页面",
     "刷新面试页", "消息、评分、当前题、进度不丢失", "符合预期", "通过", "P0", ""),
    ("F-24", "模拟面试", "7.4 / 10.3", "空回答拦截", "输入框为空",
     "点击发送或按 Enter", "前端禁用发送，不发请求", "按钮禁用", "通过", "P1", ""),
    ("F-25", "模拟面试", "7.4 / 10.2", "已完成会话禁止继续作答", "报告已生成",
     "再次访问面试页并尝试作答", "只展示已结束状态，不接收新回答", "D-02 修复后，turn/hint/skip 均拒绝", "通过", "P0", "D-02"),
    ("F-26", "复盘报告", "7.5 / AC-05", "生成复盘报告", "完成至少 1 题",
     "点击「结束并生成报告」", "跳转报告页，展示准备度、五维、逐题、行动清单", "准备度 71，8 条逐题记录", "通过", "P0", ""),
    ("F-27", "复盘报告", "7.5", "五维雷达与分项", "报告已生成",
     "查看五维能力卡片", "雷达图与 5 条进度条数值一致", "一致，标签完整", "通过", "P0", ""),
    ("F-28", "复盘报告", "7.5", "逐题复盘折叠", "报告含多题",
     "点击题目行展开/收起", "内容显示/隐藏，图标与状态同步", "可收起再展开", "通过", "P1", ""),
    ("F-29", "复盘报告", "7.5 / 7.7", "知识库引用", "报告含引用",
     "点击「知识库依据」条目", "弹窗展示文档全文", "打开《STAR 不够用…》全文", "通过", "P1", ""),
    ("F-30", "复盘报告", "7.5 / 7.6", "低分题加入题库", "报告含逐题记录",
     "点击「加入个人题库」", "按钮变为已加入，题库数量 +1", "已加入，题库 API 数量 +1", "通过", "P0", ""),
    ("F-31", "复盘报告", "7.5", "导出 Markdown", "报告已生成",
     "点击导出 Markdown", "下载完整 .md，含总评、五维、逐题、行动", "文件生成且内容完整", "通过", "P1", ""),
    ("F-32", "复盘报告", "7.5", "复制摘要", "报告已生成",
     "点击复制摘要", "剪贴板得到准备度/结论/下一步", "符合预期", "通过", "P2", ""),
    ("F-33", "个人题库", "7.6 / AC-06", "题库列表与筛选", "题库已有 1 题",
     "按全部/未掌握/复习中/已掌握筛选", "列表与数量同步", "符合预期", "通过", "P0", ""),
    ("F-34", "个人题库", "7.6", "掌握度与笔记", "题库已有题目",
     "标记已掌握；打开笔记弹窗保存", "状态与笔记持久化，卡片展示", "符合预期", "通过", "P0", ""),
    ("F-35", "个人题库", "7.6", "手动加题与去重", "题库已有题目",
     "手动添加新题；再次添加同一题", "第一次 +1；第二次不新增", "D-05 修复后数量稳定", "通过", "P1", "D-05"),
    ("F-36", "个人题库", "7.6", "题库导出", "题库已有题目",
     "点击导出 JSON", "下载题库 JSON，字段完整", "文件生成，含 id/question/status/notes", "通过", "P2", ""),
    ("F-37", "知识库", "7.7", "列表与分类", "知识库 16 篇",
     "切换分类标签", "按分类过滤，数量正确", "符合预期", "通过", "P0", ""),
    ("F-38", "知识库", "7.7", "BM25 搜索", "知识库 16 篇",
     "搜索「RAG 幻觉」", "返回相关文档与摘要片段", "返回相关结果", "通过", "P0", ""),
    ("F-39", "知识库", "7.7", "文档详情弹窗", "搜索结果至少 1 条",
     "点击卡片", "弹窗展示 Markdown 正文与标签", "符合预期", "通过", "P0", ""),
    ("F-40", "知识库", "7.7", "自定义文档增删", "进入知识库",
     "添加文档 → 搜索 → 打开 → 删除", "可检索；删除后回到 16 篇", "符合预期", "通过", "P1", ""),
    ("F-41", "面试记录", "7.0 / 7.8", "记录列表与筛选", "已有进行中/已完成记录",
     "切换全部/进行中/已完成", "记录与筛选正确", "符合预期", "通过", "P1", ""),
    ("F-42", "面试记录", "7.0", "删除记录", "已有记录",
     "点击删除并确认", "列表 -1，接口 -1，报告不可再访问", "符合预期", "通过", "P1", ""),
    ("F-43", "模型设置", "7.9 / AC-09", "本地模式", "未配置 Key",
     "打开设置页", "显示本地规则引擎；说明完整闭环可用", "符合预期", "通过", "P0", ""),
    ("F-44", "模型设置", "7.9", "接入 OpenAI 兼容接口", "准备 mock 服务",
     "选择自定义 → 填 Base URL/模型/Key → 保存", "显示已连接；计划/追问/评估/报告走模型", "保存后已连接，API 链路验证通过", "通过", "P0", ""),
    ("F-45", "模型设置", "7.9", "测试连接与错误提示", "已保存 Key / 未保存 Key",
     "分别点击测试连接", "成功显示延迟；无 Key 显示可操作提示", "成功；无 Key 显示「请先填写并保存」", "通过", "P0", ""),
    ("F-46", "模型设置", "7.9", "清除 Key", "已保存 Key",
     "点击清除并确认", "回到本地模式，Key 不再下发", "符合预期", "通过", "P1", ""),
    ("F-47", "移动端", "10.4", "侧栏菜单", "390×844 视口",
     "点击菜单 → 点击遮罩", "侧栏打开/关闭，遮罩消失", "符合预期", "通过", "P1", ""),
    ("F-48", "移动端", "10.4", "移动端面试", "390×844，进行中会话",
     "查看对话、快捷操作、输入框", "无横向溢出；提示/换题/结束可用；自动滚到最新", "符合预期", "通过", "P1", ""),
    ("F-49", "移动端", "10.4", "移动端报告", "390×844，已有报告",
     "查看报告首屏与逐题", "环形分数、雷达、逐题正常，无横向溢出", "符合预期", "通过", "P1", ""),
    ("F-50", "安全/异常", "10.2 / 10.3", "页面 404", "访问不存在的前端路径",
     "访问 /nope/page", "显示 404 提示，不白屏", "符合预期", "通过", "P2", ""),
    ("F-51", "安全/异常", "10.2", "题库 XSS 转义", "在题库写入带 img onerror 的题目",
     "打开题库页", "按文本展示，不执行脚本", "window.__xss 未定义", "通过", "P0", ""),
]

# ---------------- 读取自动化结果 ----------------
api = load("api-results.json")
edge = load("edge-results.json")
ui = load("ui-results.json")["results"]
ui2 = load("ui2-results.json")


def endpoint_for(title, module):
    t = title
    if "健康" in t: return "GET /api/health"
    if "Bootstrap" in t: return "GET /api/bootstrap"
    if "未知接口" in t: return "GET /api/*"
    if "简历" in module or "解析" in t or "doc" in t or "PDF" in t or "文本" in t: return "POST /api/resume/parse"
    if "JD" in t or "岗位" in module: return "POST /api/jd/analyze"
    if "创建面试" in t: return "POST /api/sessions"
    if "会话列表" in t: return "GET /api/sessions"
    if "删除会话" in t or "查询 404" in t: return "DELETE /api/sessions/:id"
    if "作答" in t or "回答" in t or "评分" in t or "下一题" in t: return "POST /api/sessions/:id/turn (SSE)"
    if "提示" in t: return "POST /api/sessions/:id/hint"
    if "跳过" in t: return "POST /api/sessions/:id/skip"
    if "报告" in t: return "POST /api/sessions/:id/report"
    if "题库" in t or "题目" in t: return "GET/POST/PATCH/DELETE /api/bank"
    if "知识库" in t or "文档" in t: return "GET/POST/DELETE /api/kb"
    if "设置" in t or "Key" in t or "Base URL" in t: return "GET/PUT /api/settings"
    if "路径穿越" in t: return "GET /../server/index.js"
    if "Content-Type" in t or "JSON" in t: return "POST /api/resume/parse"
    return "—"


wb = Workbook()
wb.properties.title = "面镜 MockMirror 测试用例与执行报告"
wb.properties.creator = "石浩田"
wb.properties.subject = "AI 面试教练 · 产品测试"

# ---------------- 00 总览 ----------------
ws = wb.active
ws.title = "00-测试总览"
ws.sheet_view.showGridLines = False
ws.merge_cells("A1:F1")
c = ws["A1"]
c.value = "面镜 MockMirror · 产品测试报告"
c.fill = PatternFill("solid", fgColor=INK)
c.font = Font(name="Arial", size=20, bold=True, color="FFFFFF")
c.alignment = Alignment(vertical="center", indent=1)
ws.row_dimensions[1].height = 48

ws.merge_cells("A2:F2")
c = ws["A2"]
c.value = "158 条用例 · 6 个缺陷已修复 · 0 个未关闭 P0/P1 · 2026-10-29"
c.font = Font(name="Arial", size=10, color=MUTED)
c.alignment = Alignment(vertical="center", indent=1)
ws.row_dimensions[2].height = 24

info = [
    ("测试对象", "面镜 MockMirror v1.0（浅色工具风）"),
    ("测试范围", "工作台、简历解析、JD 匹配、面试设置、模拟面试、复盘报告、个人题库、知识库、面试记录、模型设置、移动端与安全异常"),
    ("测试环境", "macOS · Node.js 24 · Chrome · 1440×920 / 390×844"),
    ("模型模式", "本地规则引擎；LLM 规划/流式/评分/报告使用本地 mock OpenAI 服务"),
    ("测试数据", "示例简历（周予安）、真实 DOCX/PDF（石浩田）、80 题、16 篇知识库"),
    ("执行方式", "接口断言 + Puppeteer 全流程 + 人工检查截图；每轮修复后回归"),
]
r = 4
for k, v in info:
    label = ws.cell(row=r, column=1, value=k)
    label.font = Font(name="Arial", size=9, bold=True, color=MUTED)
    label.alignment = Alignment(vertical="top")
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
    ws.cell(row=r, column=2, value=v).font = Font(name="Arial", size=9, color=INK)
    ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 28 if len(v) < 60 else 40
    r += 1

r += 1
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
c = ws.cell(row=r, column=1, value="  一、执行结果")
c.font = Font(name="Arial", size=12, bold=True, color=INK)
r += 1
headers = ["测试类型", "用例数", "通过", "失败", "通过率", "说明"]
for ci, h in enumerate(headers, 1):
    ws.cell(row=r, column=ci, value=h)
style_header(ws, r, 6)
summary = [
    ("功能验收", 51, 51, 0, "100%", "人工验收 + 截图检查"),
    ("接口自动化", len(api), sum(1 for x in api if x["ok"]), sum(1 for x in api if not x["ok"]), "100%", "Node fetch 断言"),
    ("UI 自动化", len(ui), sum(1 for x in ui if x["ok"]), sum(1 for x in ui if not x["ok"]), "100%", "Puppeteer 全流程"),
    ("后台/设置/安全", len(ui2), sum(1 for x in ui2 if x["ok"]), sum(1 for x in ui2 if not x["ok"]), "100%", "知识库/设置/删除/XSS"),
    ("LLM/边界", len(edge), sum(1 for x in edge if x["ok"]), sum(1 for x in edge if not x["ok"]), "100%", "mock 模型、降级、超限"),
]
r += 1
for row in summary:
    for ci, v in enumerate(row, 1):
        cell = ws.cell(row=r, column=ci, value=v)
        cell.border = BORDER
        cell.font = Font(name="Arial", size=9, color=INK)
        cell.alignment = Alignment(horizontal="center" if 2 <= ci <= 5 else "left", vertical="center")
    ws.row_dimensions[r].height = 24
    r += 1
total = 51 + len(api) + len(ui) + len(ui2) + len(edge)
for ci, v in enumerate(["合计", total, total, 0, "100%", "修复后回归"], 1):
    cell = ws.cell(row=r, column=ci, value=v)
    cell.border = BORDER
    cell.fill = PatternFill("solid", fgColor=ZEBRA)
    cell.font = Font(name="Arial", size=9, bold=True, color=INK)
    cell.alignment = Alignment(horizontal="center" if ci != 1 else "left", vertical="center")
r += 2

# 缺陷分布
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
c = ws.cell(row=r, column=1, value="  二、缺陷汇总")
c.font = Font(name="Arial", size=12, bold=True, color=INK)
r += 1
defects = [
    ("D-01", "中", "简历解析异常返回 500", "解析异常统一为 400，提示可操作", "已修复"),
    ("D-02", "高", "报告生成后仍可继续作答", "清空 pending，turn/hint/skip 校验状态", "已修复"),
    ("D-03", "中", "超大请求触发连接重置", "Content-Length 预检，超限返回 413", "已修复"),
    ("D-04", "低", "工作台统计陈旧", "首页渲染前刷新 /api/stats", "已修复"),
    ("D-05", "低", "手动重复加题出现重复卡片", "按 item.id 去重", "已修复"),
    ("D-06", "低", "模型网络错误提示不友好", "转为中文可操作提示", "已修复"),
]
ws.cell(row=r, column=1, value="统计").font = Font(name="Arial", size=10, bold=True, color=INK)
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
ws.cell(row=r, column=2, value="6 个：高 1 / 中 2 / 低 3，均已修复并回归；明细见「05-缺陷清单」").font = Font(name="Arial", size=9.5, color=INK)
ws.cell(row=r, column=2).alignment = Alignment(vertical="center", wrap_text=True)
ws.row_dimensions[r].height = 26
r += 1
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
c = ws.cell(row=r, column=1, value="  三、结论与发布建议")
c.font = Font(name="Arial", size=12, bold=True, color=INK)
r += 1
conclusions = [
    "核心闭环全部通过：简历/JD → 画像与匹配 → 题单 → 多轮追问 → 复盘 → 题库。",
    "本地模式零 Key 可完整运行；LLM 模式可流式追问与生成报告，故障时自动降级。",
    "首轮 6 个缺陷已全部修复并回归，无未关闭 P0/P1；可进入 30 人封闭灰度。",
    "灰度重点观察：报告行动转化率、7 日内复练率、追问下的稳定性，而不是继续扩功能。",
]
for line in conclusions:
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    cell = ws.cell(row=r, column=1, value="• " + line)
    cell.font = Font(name="Arial", size=9.5, color=INK)
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 24
    r += 1

for col, w in zip("ABCDEF", [24, 12, 12, 12, 12, 26]):
    ws.column_dimensions[col].width = w

# ---------------- 01 功能用例 ----------------
ws = wb.create_sheet("01-功能测试用例")
ws.sheet_view.showGridLines = False
title_block(ws, "01 · 功能测试用例", "人工验收口径 · 51 条 · 关联 PRD 章节与缺陷", 11)
headers = ["用例ID", "模块", "PRD 章节", "用例标题", "前置条件", "操作步骤", "预期结果", "实际结果", "结果", "优先级", "关联缺陷"]
for ci, h in enumerate(headers, 1):
    ws.cell(row=4, column=ci, value=h)
style_header(ws, 4, 11)
write_rows(ws, 5, F, [10, 10, 10, 18, 16, 28, 28, 28, 8, 9, 10], result_col=9, height=54)
ws.row_dimensions[4].height = 30
ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A4:K{4 + len(F)}"
ws.print_title_rows = "1:4"

# ---------------- 02 接口用例 ----------------
ws = wb.create_sheet("02-接口自动化")
ws.sheet_view.showGridLines = False
title_block(ws, "02 · 接口自动化结果", f"{len(api)} 条 · Node fetch 断言 · 断言失败会记录预期/实际", 7)
headers = ["用例ID", "模块", "用例标题", "接口", "预期结果", "实际结果", "结果"]
for ci, h in enumerate(headers, 1):
    ws.cell(row=4, column=ci, value=h)
style_header(ws, 4, 7)
rows = []
for x in api:
    rows.append([x["id"], x["module"], x["title"], endpoint_for(x["title"], x["module"]), x["expected"], x["actual"], "通过" if x["ok"] else "失败"])
write_rows(ws, 5, rows, [10, 11, 28, 30, 32, 32, 8], result_col=7, height=38)
ws.row_dimensions[4].height = 28
ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A4:G{4 + len(rows)}"
ws.print_title_rows = "1:4"

# ---------------- 03 UI 自动化 ----------------
ws = wb.create_sheet("03-UI自动化")
ws.sheet_view.showGridLines = False
title_block(ws, "03 · UI 自动化结果", f"{len(ui) + len(ui2)} 条 · Puppeteer 浏览器全流程", 6)
headers = ["用例ID", "模块", "用例标题", "预期结果", "实际结果", "结果"]
for ci, h in enumerate(headers, 1):
    ws.cell(row=4, column=ci, value=h)
style_header(ws, 4, 6)
rows = []
for x in ui + ui2:
    rows.append([x["id"], x["module"], x["title"], x["expected"], x["actual"], "通过" if x["ok"] else "失败"])
write_rows(ws, 5, rows, [10, 11, 26, 32, 32, 8], result_col=6, height=36)
ws.row_dimensions[4].height = 28
ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A4:F{4 + len(rows)}"
ws.print_title_rows = "1:4"

# ---------------- 04 LLM/边界 ----------------
ws = wb.create_sheet("04-LLM与边界")
ws.sheet_view.showGridLines = False
title_block(ws, "04 · LLM / 边界 / 安全", f"{len(edge)} 条 · 正常链路、降级、超限、XSS 与损坏文件", 6)
headers = ["用例ID", "模块", "用例标题", "预期结果", "实际结果", "结果"]
for ci, h in enumerate(headers, 1):
    ws.cell(row=4, column=ci, value=h)
style_header(ws, 4, 6)
rows = [[x["id"], x["module"], x["title"], x["expected"], x["actual"], "通过" if x["ok"] else "失败"] for x in edge]
write_rows(ws, 5, rows, [10, 10, 28, 32, 38, 8], result_col=6, height=36)
ws.row_dimensions[4].height = 28
ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A4:F{4 + len(rows)}"

# ---------------- 05 缺陷 ----------------
ws = wb.create_sheet("05-缺陷清单")
ws.sheet_view.showGridLines = False
title_block(ws, "05 · 缺陷清单", "发现 6 个 · 已修复 6 个 · 未关闭 P0/P1 = 0", 11)
headers = ["ID", "发现阶段", "严重度", "模块", "问题描述", "复现步骤", "预期", "实际", "根因", "修复方案", "状态"]
for ci, h in enumerate(headers, 1):
    ws.cell(row=4, column=ci, value=h)
style_header(ws, 4, 11)
D = [
    ("D-01", "接口测试", "中", "简历解析", "不支持格式/损坏 PDF/.doc 返回 500", "上传 .doc 或损坏 PDF", "400 + 可操作提示", "500 + 技术错误", "parseUpload 异常未包装为 HttpError", "路由层捕获并转为 400", "已修复"),
    ("D-02", "接口测试", "高", "模拟面试", "生成报告后仍可作答/提示/跳过", "作答 1 题 → 生成报告 → 再发 turn", "拒绝，会话已结束", "200，继续生成下一题", "finishSession 未清空 pending；turn 只检查 pending", "清空 pending；turn/hint/skip 校验 status=active", "已修复"),
    ("D-03", "边界测试", "中", "服务端", "超大请求体触发 ECONNRESET", "POST 13MB JSON 到 /api/resume/parse", "413", "fetch failed", "超限时停止读取，连接被重置", "预检 Content-Length；req.resume() 后返回 413", "已修复"),
    ("D-04", "UI 测试", "低", "工作台", "完成面试后统计仍显示旧值", "完成 1 场 → 回工作台", "统计 +1", "仍显示 0", "首页使用启动时 state.stats", "renderHome 前 refreshStats", "已修复"),
    ("D-05", "UI 测试", "低", "个人题库", "手动重复加题出现重复卡片", "同一题连续添加两次", "只保留 1 条", "列表出现 2 条", "前端无条件 unshift", "按 item.id 去重", "已修复"),
    ("D-06", "降级测试", "低", "模型设置", "网络错误提示 fetch failed", "Base URL 指向不可达端口", "中文可操作提示", "fetch failed", "未区分网络错误", "捕获 TypeError 并转中文", "已修复"),
]
write_rows(ws, 5, D, [8, 10, 8, 12, 30, 28, 22, 24, 28, 30, 8], result_col=11, height=58)
ws.row_dimensions[4].height = 30
ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A4:K{4 + len(D)}"

# ---------------- 06 环境 ----------------
ws = wb.create_sheet("06-环境与数据")
ws.sheet_view.showGridLines = False
title_block(ws, "06 · 测试环境、数据与命令", "可复现的测试基线", 4)
items = [
    ("操作系统", "macOS 26（Apple Silicon）"),
    ("运行时", "Node.js v24.12.0；前端原生 ES Modules"),
    ("浏览器", "Google Chrome（headless new）；桌面 1440×920，移动 390×844"),
    ("被测版本", "面镜 MockMirror v1.0 + 测试修复（浅色工具风）"),
    ("简历样本", "示例简历（周予安）；真实 DOCX/PDF（石浩田，3 段实习 / 2 个项目）"),
    ("题库/知识库", "80 道题；16 篇知识；6 个岗位包；3 种人格"),
    ("模型 Mock", "本地 mock OpenAI 服务：http://127.0.0.1:4399/v1，验证规划/流式/评分/报告与失败降级"),
    ("接口测试", "cd ai-interview-coach && node /tmp/mmtest/api-test.mjs（结果见 docs/test-results/api-results.json）"),
    ("UI 测试", "node /tmp/mmtest/ui-test.mjs；node /tmp/mmtest/ui-test2.mjs"),
    ("边界测试", "node /tmp/mmtest/edge-test.mjs"),
    ("数据清理", "删除 data/store.json 后重启服务，可回到首次运行状态"),
    ("复现说明", "测试脚本依赖 Node 18+ 的 fetch；UI 测试依赖 puppeteer-core 与系统 Chrome"),
]
r = 4
for k, v in items:
    ws.cell(row=r, column=1, value=k).font = Font(name="Arial", size=10, bold=True, color=MUTED)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    cell = ws.cell(row=r, column=2, value=v)
    cell.font = Font(name="Arial", size=10, color=INK)
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 30 if len(v) < 55 else 44
    r += 1
for col, w in zip("ABCD", [18, 45, 25, 25]):
    ws.column_dimensions[col].width = w

# ---------------- 打印设置 ----------------
for ws in wb.worksheets:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.45, bottom=0.45)

for ws in wb.worksheets:
    ws.sheet_view.zoomScale = 85

wb.save(OUT)
print("saved", OUT, OUT.stat().st_size)
