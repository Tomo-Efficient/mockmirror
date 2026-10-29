#!/usr/bin/env python3
"""把 docs/PRD.md 转成可交付的 Word 文档（保留标题、表格、列表、引用与代码块）。"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "docs" / "PRD.md"
OUT = ROOT / "docs" / "面镜MockMirror-产品需求文档-v1.1.docx"

BODY_FONT = "等线"
LATIN_FONT = "Arial"
MONO_FONT = "Courier New"
INK = RGBColor(0x11, 0x18, 0x27)
MUTED = RGBColor(0x4B, 0x55, 0x63)
ACCENT = RGBColor(0x1F, 0x29, 0x37)


def set_east_asia(style_or_run, name=BODY_FONT):
    rpr = style_or_run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:ascii"), LATIN_FONT)
    rfonts.set(qn("w:hAnsi"), LATIN_FONT)


def shade(element, color):
    ppr = element.get_or_add_pPr() if hasattr(element, "get_or_add_pPr") else element
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    ppr.append(shd)


def cell_shade(cell, color):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcpr.append(shd)


def set_cell_margins(cell, top=50, bottom=50, left=80, right=80):
    tcpr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for tag, value in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tcpr.append(mar)


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    trpr.append(node)


def no_split(row):
    trpr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    node.set(qn("w:val"), "true")
    trpr.append(node)


def add_page_field(paragraph, field):
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t")
    txt.text = "1"
    fld3 = OxmlElement("w:fldChar")
    fld3.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)
    run._r.append(txt)
    run._r.append(fld3)


def bottom_border(paragraph, color="D1D5DB", size=4):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    ppr.append(pbdr)


def add_inline(paragraph, text, size=None, color=None, bold_all=False):
    """解析 **bold**、`code`、[text](url)。"""
    pattern = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            run.bold = bold_all
            if size: run.font.size = Pt(size)
            if color: run.font.color.rgb = color
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = MONO_FONT
            run.font.size = Pt((size or 10.5) - 1)
            run.font.color.rgb = RGBColor(0xB9, 0x1C, 0x1C)
        else:
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            run = paragraph.add_run(m.group(1))
            run.font.color.rgb = RGBColor(0x1D, 0x4E, 0xD8)
            run.underline = True
            if m.group(2).startswith("http"):
                run2 = paragraph.add_run(" " + m.group(2))
                run2.font.size = Pt(8)
                run2.font.color.rgb = MUTED
        if size: run.font.size = Pt(size)
        if color and not token.startswith("`"):
            run.font.color.rgb = color
        if bold_all and not token.startswith("**"):
            run.bold = True
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        run.bold = bold_all
        if size: run.font.size = Pt(size)
        if color: run.font.color.rgb = color


def add_cover(doc):
    section = doc.sections[0]
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run("PRODUCT REQUIREMENTS DOCUMENT")
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = MUTED
    run.font.name = LATIN_FONT

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    run = p.add_run("面镜 MockMirror")
    run.font.size = Pt(38)
    run.font.bold = True
    run.font.color.rgb = INK
    set_east_asia(run, "等线 Light")

    p = doc.add_paragraph()
    run = p.add_run("AI 面试教练 · 产品需求文档")
    run.font.size = Pt(18)
    run.font.color.rgb = ACCENT

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    run = p.add_run("把每一次面试，提前演练一遍。")
    run.italic = True
    run.font.size = Pt(12)
    run.font.color.rgb = MUTED

    p = doc.add_paragraph()
    bottom_border(p, "111827", 8)
    p.paragraph_format.space_before = Pt(20)

    for _ in range(3):
        doc.add_paragraph()

    table = doc.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta = [
        ("文档版本", "v1.1（含 v1.0 验收结果）"),
        ("文档日期", "2026-10-29"),
        ("作者", "石浩田 · 香港城市大学统计学硕士 · AI / 数据产品方向"),
        ("文档状态", "核心闭环已实现，158 条用例通过，进入灰度准备"),
        ("适用范围", "产品验收、灰度发布、后续迭代与作品集说明"),
    ]
    for i, (key, value) in enumerate(meta):
        c0, c1 = table.rows[i].cells
        c0.width = Cm(3.2)
        c1.width = Cm(12.2)
        cell_shade(c0, "F3F4F6")
        set_cell_margins(c0); set_cell_margins(c1)
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(key)
        r0.bold = True
        r0.font.size = Pt(9.5)
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(value)
        r1.font.size = Pt(9.5)

    doc.add_page_break()


def add_revision_and_toc(doc, lines):
    p = doc.add_paragraph()
    run = p.add_run("修订记录")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = INK
    p.paragraph_format.space_after = Pt(6)

    rows = [
        ("v1.0", "2026-10-24", "初版：背景、用户、指标、流程、9 个功能模块、数据模型、风险与路线图"),
        ("v1.1", "2026-10-29", "补充 158 条测试验收结果与 6 个缺陷修复；视觉规范切换为浅色工具风"),
    ]
    table = doc.add_table(rows=1 + len(rows), cols=3)
    table.style = "Table Grid"
    headers = ("版本", "日期", "变更内容")
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell_shade(cell, "E5E7EB")
        set_cell_margins(cell)
        r = cell.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(9)
    for ri, row in enumerate(rows, 1):
        for ci, value in enumerate(row):
            cell = table.rows[ri].cells[ci]
            set_cell_margins(cell)
            r = cell.paragraphs[0].add_run(value)
            r.font.size = Pt(9)
    table.columns[0].width = Cm(2.2)
    table.columns[1].width = Cm(3.2)
    table.columns[2].width = Cm(10.2)
    doc.add_paragraph()

    p = doc.add_paragraph()
    run = p.add_run("目录")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = INK
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)

    for line in lines:
        m = re.match(r"##\s+(\d+)\.\s+(.+)", line.strip())
        if not m: continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(15.5), WD_TAB_ALIGNMENT.RIGHT, 1)
        r = p.add_run(f"{m.group(1)}  {m.group(2)}")
        r.font.size = Pt(10)
        r.font.color.rgb = ACCENT
    doc.add_page_break()


def add_table(doc, header, rows):
    cols = len(header)
    table = doc.add_table(rows=1 + len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    size = 8.2 if cols >= 6 else 9
    for ci, text in enumerate(header):
        cell = table.rows[0].cells[ci]
        cell_shade(cell, "F3F4F6")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_inline(p, text, size=size, bold_all=True)
    repeat_header(table.rows[0])
    no_split(table.rows[0])
    for ri, row in enumerate(rows, 1):
        for ci in range(cols):
            cell = table.rows[ri].cells[ci]
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, row[ci] if ci < len(row) else "", size=size)
        no_split(table.rows[ri])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_code_block(doc, code, lang=""):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.3)
    p.paragraph_format.right_indent = Cm(0.3)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    shade(p._p, "F6F8FA")
    for i, line in enumerate(code.split("\n")):
        run = p.add_run(line)
        run.font.name = MONO_FONT
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
        if i < len(code.split("\n")) - 1:
            run.add_break()


def add_blockquote(doc, quote):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.right_indent = Cm(0.2)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    shade(p._p, "F9FAFB")
    ppr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single"); left.set(qn("w:sz"), "12")
    left.set(qn("w:space"), "8"); left.set(qn("w:color"), "D1D5DB")
    pbdr.append(left); ppr.append(pbdr)
    add_inline(p, quote, size=10, color=MUTED)


def build_body(doc, lines):
    i = 0
    n = len(lines)
    first_h2 = True
    while i < n:
        raw = lines[i]
        line = raw.rstrip("\n")
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("```"):
            lang = stripped[3:].strip()
            code = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                code.append(lines[i].rstrip("\n"))
                i += 1
            i += 1
            add_code_block(doc, "\n".join(code), lang)
            continue
        if stripped.startswith("|") and i + 1 < n and re.match(r"^\|[\s:|-]+\|?", lines[i + 1].strip()):
            header = [c.strip() for c in stripped.strip("|").split("|")]
            i += 2
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            add_table(doc, header, rows)
            continue
        if stripped.startswith(">"):
            buf = []
            while i < n and lines[i].strip().startswith(">"):
                buf.append(re.sub(r"^\s*>\s?", "", lines[i]).rstrip())
                i += 1
            add_blockquote(doc, " ".join(buf))
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            title = m.group(2).strip()
            if level == 1:
                i += 1
                continue
            if level == 2:
                p = doc.add_paragraph(style="Heading 1")
                if not first_h2:
                    p.paragraph_format.page_break_before = True
                first_h2 = False
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(10)
                add_inline(p, title, size=17, bold_all=True)
                bottom_border(p, "D1D5DB", 6)
            elif level == 3:
                p = doc.add_paragraph(style="Heading 2")
                p.paragraph_format.space_before = Pt(14)
                p.paragraph_format.space_after = Pt(6)
                add_inline(p, title, size=13, bold_all=True)
            else:
                p = doc.add_paragraph(style="Heading 3")
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(4)
                add_inline(p, title, size=11, bold_all=True)
            i += 1
            continue
        if re.match(r"^\s*[-*]\s+", raw):
            indent = len(raw) - len(raw.lstrip(" "))
            ordered = False
            style = "List Bullet" if indent < 2 else "List Bullet 2"
            text = re.sub(r"^\s*[-*]\s+", "", raw)
            p = doc.add_paragraph(style=style)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, text, size=10.5)
            i += 1
            continue
        if re.match(r"^\s*\d+\.\s+", raw):
            idx = 1
            while i < n and re.match(r"^\s*\d+\.\s+", lines[i]):
                text = re.sub(r"^\s*\d+\.\s+", "", lines[i]).rstrip()
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.75)
                p.paragraph_format.first_line_indent = Cm(-0.45)
                p.paragraph_format.space_after = Pt(2)
                run = p.add_run(f"{idx}.  ")
                run.bold = True
                add_inline(p, text, size=10.5)
                idx += 1
                i += 1
            continue
        if re.match(r"^\s*---+\s*$", stripped):
            p = doc.add_paragraph()
            bottom_border(p, "E5E7EB", 4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            i += 1
            continue
        # paragraph: collect until blank/special
        buf = [stripped]
        i += 1
        while i < n:
            nxt = lines[i].strip()
            if not nxt or nxt.startswith(("#", "|", ">", "```", "---")) or re.match(r"^([-*]|\d+\.)\s+", lines[i]):
                break
            buf.append(nxt)
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        add_inline(p, "".join(buf) if all(len(x) < 60 for x in buf) else " ".join(buf), size=10.5)


def add_header_footer(doc, title):
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(title)
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED
    bottom_border(p, "E5E7EB", 4)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("面镜 MockMirror · 内部文档 · ")
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED
    run2 = p.add_run("第 ")
    run2.font.size = Pt(8); run2.font.color.rgb = MUTED
    add_page_field(p, "PAGE")
    run3 = p.add_run(" 页 / 共 ")
    run3.font.size = Pt(8); run3.font.color.rgb = MUTED
    add_page_field(p, "NUMPAGES")
    run4 = p.add_run(" 页")
    run4.font.size = Pt(8); run4.font.color.rgb = MUTED


def main():
    text = SRC.read_text(encoding="utf-8")
    lines = text.splitlines()

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = LATIN_FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    set_east_asia(normal)
    normal.paragraph_format.line_spacing = 1.55
    normal.paragraph_format.space_after = Pt(5)

    for name, size, color in (("Heading 1", 17, INK), ("Heading 2", 13, ACCENT), ("Heading 3", 11, ACCENT)):
        st = styles[name]
        st.font.name = LATIN_FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        set_east_asia(st)

    doc.core_properties.title = "面镜 MockMirror 产品需求文档"
    doc.core_properties.author = "石浩田"
    doc.core_properties.subject = "AI 面试教练产品需求与验收"
    doc.core_properties.comments = "v1.1 · 158 条用例通过"

    add_cover(doc)
    add_revision_and_toc(doc, lines)
    build_body(doc, lines)
    add_header_footer(doc, "面镜 MockMirror · 产品需求文档 v1.1")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"saved {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
